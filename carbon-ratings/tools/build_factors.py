#!/usr/bin/env python3
"""
build_factors.py -- reproducible build of data/factors.csv for the
India-EU Carbon Ratings engine (SPEC sections 4-5).

HOUSE RULE: no value in factors.csv may come from memory. Every value_primary
is parsed out of a file downloaded by this script into data/raw/. If a value
cannot be extracted from a source, the row is emitted with an empty
value_primary and flags=TODO -- never guessed.

Run:      python tools/build_factors.py
Re-run:   cached files under data/raw/ are reused; pass --refresh to re-download.

Requires: pandas openpyxl requests pdfplumber
"""

from __future__ import annotations

import argparse
import csv
import datetime as _dt
import re
import sys
from pathlib import Path

import openpyxl
import pdfplumber
import requests

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "factors.csv"

RETRIEVED = "2026-08-27"  # date the cached copies in data/raw/ were downloaded

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}

SOURCES = {
    # local filename -> canonical URL
    "GHGP_Cross_Sector_Tools_V2.0.xlsx":
        "https://ghgprotocol.org/sites/default/files/2024-05/"
        "Emission_Factors_for_Cross_Sector_Tools_V2.0_0.xlsx",
    "DEFRA_2026_full_set.xlsx":
        "https://assets.publishing.service.gov.uk/media/6a29392bade52dc0882218a8/"
        "ghg-conversion-factors-2026-full-set.xlsx",
    "DEFRA_2026_flat_format.xlsx":
        "https://assets.publishing.service.gov.uk/media/6a6c9748862aaf18d9c62ac9/"
        "ghg-conversion-factors-2026-flat-format-revised.xlsx",
    "DEFRA_2026_methodology.pdf":
        "https://assets.publishing.service.gov.uk/media/6a2940543b15d05a7ce3202e/"
        "2026-GHG-conversion-factors-methodology-report.pdf",
    "CEA_User_Guide_V21.0.pdf":
        "https://cea.nic.in/wp-content/uploads/baseline/2025/12/User_Guide_V_21.0.pdf",
    "IPCC_2006_V3_Ch2_Mineral.pdf":
        "https://www.ipcc-nggip.iges.or.jp/public/2006gl/pdf/3_Volume3/"
        "V3_2_Ch2_Mineral_Industry.pdf",
    "IPCC_2006_V3_Ch3_Chemical.pdf":
        "https://www.ipcc-nggip.iges.or.jp/public/2006gl/pdf/3_Volume3/"
        "V3_3_Ch3_Chemical_Industry.pdf",
    "IPCC_2006_V3_Ch4_Metal.pdf":
        "https://www.ipcc-nggip.iges.or.jp/public/2006gl/pdf/3_Volume3/"
        "V3_4_Ch4_Metal_Industry.pdf",
    "IPCC_AR6_WGI_Chapter07.pdf":
        "https://www.ipcc.ch/report/ar6/wg1/downloads/report/IPCC_AR6_WGI_Chapter07.pdf",
    "IPCC_AR6_WGI_Chapter07_SM.pdf":
        "https://www.ipcc.ch/report/ar6/wg1/downloads/report/IPCC_AR6_WGI_Chapter07_SM.pdf",
    "EPA_SNAP_refrigerant_blend_compositions.html":
        "https://www.epa.gov/snap/compositions-refrigerant-blends",
}

# Source citation strings used in the CSV
S_GHGP = "GHG Protocol Cross-Sector Tools EF workbook v2.0"
S_DEFRA = "UK DESNZ/DEFRA GHG conversion factors 2026 (full set)"
S_DEFRA_METH = "UK DESNZ/DEFRA 2026 conversion factors methodology paper"
S_CEA = "CEA CO2 Baseline Database for the Indian Power Sector v21.0"
S_IPCC2 = "IPCC 2006 GL Vol.3 Ch.2 (Mineral Industry)"
S_IPCC3 = "IPCC 2006 GL Vol.3 Ch.3 (Chemical Industry)"
S_IPCC4 = "IPCC 2006 GL Vol.3 Ch.4 (Metal Industry)"
S_AR6 = "IPCC AR6 WGI Ch.7 Table 7.15"
S_AR6SM = "IPCC AR6 WGI Ch.7 Supplementary Material Table 7.SM.7"
S_EPA_SNAP = "US EPA SNAP, Compositions of Refrigerant Blends"

V_GHGP = "v2.0 (Mar 2024; IPCC 2006 basis)"
V_DEFRA = "2026 (pub. 11 Jun 2026, last upd. 31 Jul 2026)"
V_CEA = "v21.0 (Nov 2025), FY 2024-25"
V_IPCC = "2006 Guidelines"
V_AR6 = "AR6 (2021) GWP-100"


# --------------------------------------------------------------------------
# fetch
# --------------------------------------------------------------------------
def fetch_all(refresh: bool = False) -> None:
    RAW.mkdir(parents=True, exist_ok=True)
    for name, url in SOURCES.items():
        dest = RAW / name
        if dest.exists() and not refresh:
            print(f"  cached  {name} ({dest.stat().st_size:,} B)")
            continue
        print(f"  GET     {name} <- {url}")
        r = requests.get(url, headers=UA, timeout=600)
        r.raise_for_status()
        dest.write_bytes(r.content)
        print(f"          {len(r.content):,} B written")


# --------------------------------------------------------------------------
# parsers
# --------------------------------------------------------------------------
def parse_ghgp_fuels() -> dict:
    """GHGP 'Stationary Combustion' Tables 1-3.

    Layout (verified against the workbook): col B category, col C fuel,
    D NCV TJ/Gg, E energy basis, F mass basis, G liquid density kg/L,
    H gas density kg/m3, I per-litre, J per-m3.
    CO2 block rows 6-59, CH4 block 71-124, N2O block 136-189 -- same fuel
    order in all three, which we assert on the NCV column.
    """
    wb = openpyxl.load_workbook(RAW / "GHGP_Cross_Sector_Tools_V2.0.xlsx",
                                data_only=True)
    ws = wb["Stationary Combustion"]
    blocks = {"CO2": 6, "CH4": 71, "N2O": 136}
    n = 54  # 54 fuels per table

    def read_block(start):
        out = []
        for i in range(n):
            r = start + i
            out.append({
                "fuel": (ws.cell(r, 3).value or "").strip(),
                "ncv": ws.cell(r, 4).value,
                "per_tj": ws.cell(r, 5).value,
                "per_tonne": ws.cell(r, 6).value,
                "dens_l": ws.cell(r, 7).value,
                "dens_m3": ws.cell(r, 8).value,
                "per_litre": ws.cell(r, 9).value,
                "per_m3": ws.cell(r, 10).value,
            })
        return out

    co2, ch4, n2o = (read_block(blocks[g]) for g in ("CO2", "CH4", "N2O"))
    for a, b, c in zip(co2, ch4, n2o):
        assert a["ncv"] == b["ncv"] == c["ncv"], f"GHGP block misalignment at {a['fuel']}"

    fuels = {}
    for i, row in enumerate(co2):
        key = _norm_fuel(row["fuel"])
        fuels[key] = {
            "label": row["fuel"], "ncv": row["ncv"],
            "dens_l": row["dens_l"], "dens_m3": row["dens_m3"],
            "CO2": co2[i], "CH4": ch4[i], "N2O": n2o[i],
        }
    return fuels


def _norm_fuel(s: str) -> str:
    """Strip the trailing footnote marker: 'Residual fuel oil3' -> 'residual fuel oil'.

    Do NOT split on '/': the CO2 block names the diesel row 'Gas/Diesel oil', and
    splitting there silently collapses it to 'gas'.
    """
    return re.sub(r"\d+$", "", (s or "").strip()).strip().lower()


def parse_defra_fuels() -> dict:
    """DEFRA 2026 'Fuels' sheet -> {(fuel_lower, unit_lower): {...}}.

    Columns: A Activity, B Fuel, C Unit, D kgCO2e, E kgCO2e of CO2,
    F kgCO2e of CH4, G kgCO2e of N2O. Fuel name is merged across its unit
    rows, so we forward-fill it.
    """
    wb = openpyxl.load_workbook(RAW / "DEFRA_2026_full_set.xlsx", data_only=True)
    ws = wb["Fuels"]
    out, fuel = {}, None
    for r in range(23, 149):
        b = ws.cell(r, 2).value
        unit = ws.cell(r, 3).value
        if isinstance(b, str) and b.strip():
            fuel = b.strip()
        if not (fuel and isinstance(unit, str) and unit.strip()):
            continue
        co2e, co2, ch4, n2o = (ws.cell(r, c).value for c in (4, 5, 6, 7))
        if co2 is None:
            continue
        out[(fuel.lower(), unit.strip().lower())] = {
            "co2e": co2e, "co2": co2, "ch4_co2e": ch4, "n2o_co2e": n2o,
        }
    assert ("diesel (100% mineral diesel)", "litres") in out, "DEFRA Fuels parse failed"
    return out


def parse_defra_gwp() -> dict:
    """DEFRA 2026 'Refrigerant & other' -> {name_lower: total kgCO2e/kg} (AR5)."""
    wb = openpyxl.load_workbook(RAW / "DEFRA_2026_full_set.xlsx", data_only=True)
    ws = wb["Refrigerant & other"]
    out = {}
    for r in range(20, 203):
        name = ws.cell(r, 2).value
        total = ws.cell(r, 6).value
        if isinstance(name, str) and isinstance(total, (int, float)):
            out[name.strip().lower()] = float(total)
    assert out.get("methane") == 28, "DEFRA GWP parse failed (expected AR5 CH4=28)"
    return out


def parse_cea() -> dict:
    """CEA v21.0 User Guide: Table S (FY2024-25) + Annexure-I time series."""
    with pdfplumber.open(RAW / "CEA_User_Guide_V21.0.pdf") as pdf:
        txt = "\n".join((p.extract_text() or "") for p in pdf.pages)

    m = re.search(r"Average\s+OM\s+BM\s+CM\s*\n\s*"
                  r"([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)", txt)
    if not m:
        raise RuntimeError("CEA Table S not found")
    res = {"avg": float(m.group(1)), "om": float(m.group(2)),
           "bm": float(m.group(3)), "cm": float(m.group(4))}

    # Annexure-I: prior-vintage weighted average (FY 2023-24) for cross-check
    m2 = re.search(r"2023-24\s+[\d.]+\s+[\d.]+\s+[\d.]+\s+[\d.]+\s+([\d.]+)", txt)
    res["prev_fy_avg"] = float(m2.group(1)) if m2 else None
    return res


def parse_ipcc_clinker() -> dict:
    with pdfplumber.open(RAW / "IPCC_2006_V3_Ch2_Mineral.pdf") as pdf:
        txt = "\n".join((p.extract_text() or "") for p in pdf.pages)
    m = re.search(r"EF\s*=\s*0\.51\s*\W?\s*1\.02\s*\(CKD\s*correction\)\s*=\s*"
                  r"([\d.]+)\s*tonnes\s*CO", txt)
    if not m:
        raise RuntimeError("IPCC Eq 2.4 clinker EF not found")
    out = {"clinker_tier1": float(m.group(1)), "clinker_base": 0.51}
    # Table 2.1 carbonate EFs (tonnes CO2 / tonne carbonate)
    for name, pat in (("caco3", r"Calcite\*{0,3}\s*or\s*aragonite\s+100\.0869\s+([\d.]+)"),
                      ("dolomite", r"Dolomite\*{0,3}\s+184\.4008\s+([\d.]+)")):
        mm = re.search(pat, txt)
        if not mm:
            raise RuntimeError(f"IPCC Table 2.1 row {name} not found")
        out[name] = float(mm.group(1))
    return out


def parse_ipcc_chemical() -> dict:
    with pdfplumber.open(RAW / "IPCC_2006_V3_Ch3_Chemical.pdf") as pdf:
        txt = "\n".join((p.extract_text() or "") for p in pdf.pages)
    out = {}
    # Table 3.3 default N2O factors for nitric acid production
    # NB: pdfplumber renders the subscript in "N2O" inline here, but that is not
    # guaranteed across the chapter, so match N, an optional subscript char, then O.
    _n2o = r"N\s*2?\s*O"
    pats = {
        "nscr": rf"Plants with NSCRa?\s*\(all processes\)\s+([\d.]+)\s*kg\s*{_n2o}",
        "destruction": rf"tailgas\s+{_n2o}\s+destruction\s+([\d.]+)\s*kg\s*{_n2o}",
        "atmospheric": rf"Atmospheric pressure plants \(low pressure\)\s+([\d.]+)\s*kg\s*{_n2o}",
        "medium": rf"Medium pressure combustion plants\s+([\d.]+)\s*kg\s*{_n2o}",
        "high": rf"High pressure plants\s+([\d.]+)\s*kg\s*{_n2o}",
    }
    for k, p in pats.items():
        m = re.search(p, txt)
        if not m:
            raise RuntimeError(f"IPCC Table 3.3 row '{k}' not found")
        out["hno3_" + k] = float(m.group(1))
    # Urea: CO2 consumed per tonne urea produced. The sentence is split by a line
    # of stray subscript digits ("...per\n3 2 2\ntonne of urea produced"), so allow
    # a short run of any characters between the two halves.
    m = re.search(r"([\d.]+)\s*tonnes of CO\s*2?\s*are required per"
                  r"[\s\S]{0,30}?tonne of urea produced", txt)
    if not m:
        raise RuntimeError("IPCC urea CO2 uptake not found")
    out["urea_uptake"] = float(m.group(1))
    return out


def parse_ipcc_metal() -> dict:
    with pdfplumber.open(RAW / "IPCC_2006_V3_Ch4_Metal.pdf") as pdf:
        txt = "\n".join((p.extract_text() or "") for p in pdf.pages)
    out = {}
    # Table 4.1 Tier 1 CO2 EFs (tonne CO2 / tonne product). The label cells wrap
    # across lines with stray subscripts, so read the ruled table rather than
    # regexing the reflowed text.
    out.update(_parse_table_4_1())
    # Table 4.3 material carbon contents (kg C / kg)
    m = re.search(r"EAF Carbon Electrodes\d?\s+([\d.]+)", txt)
    if not m:
        raise RuntimeError("IPCC Table 4.3 EAF electrode carbon content not found")
    out["eaf_electrode_c"] = float(m.group(1))
    # Table 4.10 anode/paste CO2 (tonne CO2 / tonne Al)
    for k, p in (("prebake", r"Prebake\d?\s+([\d.]+)\s+10"),
                 ("soderberg", r"S.derberg\s+([\d.]+)\s+10")):
        mm = re.search(p, txt)
        if not mm:
            raise RuntimeError(f"IPCC Table 4.10 row '{k}' not found")
        out["al_anode_" + k] = float(mm.group(1))
    # Table 4.15 Tier 1 PFC defaults (kg / tonne Al), 4 technologies x 2 gases
    m = re.search(r"CWPB\s*\n\s*SWPB\s*\n\s*VSS\s*\n\s*HSS\s+"
                  r"([\d.]+)\s*\n\s*([\d.]+)\s*\n\s*([\d.]+)\s*\n\s*([\d.]+)\s+"
                  r"[-\d/+]+\s*\n\s*[-\d/+]+\s*\n\s*[-\d/+]+\s*\n\s*[-\d/+]+\s+"
                  r"([\d.]+)\s*\n\s*([\d.]+)\s*\n\s*([\d.]+)\s*\n\s*([\d.]+)", txt)
    if m:
        g = [float(x) for x in m.groups()]
        out["pfc"] = {"CWPB": (g[0], g[4]), "SWPB": (g[1], g[5]),
                      "VSS": (g[2], g[6]), "HSS": (g[3], g[7])}
    else:
        # fall back to the table-extraction route
        out["pfc"] = _parse_pfc_via_tables()
    return out


_T41_LABELS = {
    "sinter production": "steel_sinter",
    "coke oven": "steel_coke_oven",
    "iron production": "steel_bf_iron",
    "direct reduced iron production": "steel_dri",
    "basic oxygen furnace": "steel_bof",
    "electric arc furnace": "steel_eaf",
}


def _parse_table_4_1() -> dict:
    """IPCC Vol.3 Ch.4 Table 4.1 via ruled-table extraction."""
    out = {}
    with pdfplumber.open(RAW / "IPCC_2006_V3_Ch4_Metal.pdf") as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            if "TABLE 4.1" not in t or "TIER 1 DEFAULT" not in t:
                continue
            for tb in (page.extract_tables() or []):
                for row in tb:
                    if not row or len(row) < 2 or not row[0] or not row[1]:
                        continue
                    label = " ".join(row[0].split()).lower()
                    try:
                        val = float(row[1].strip())
                    except ValueError:
                        continue
                    for prefix, key in _T41_LABELS.items():
                        if label.startswith(prefix) and key not in out:
                            out[key] = val
    missing = set(_T41_LABELS.values()) - set(out)
    if missing:
        raise RuntimeError(f"IPCC Table 4.1 rows not found: {sorted(missing)}")
    return out


def _parse_pfc_via_tables() -> dict:
    with pdfplumber.open(RAW / "IPCC_2006_V3_Ch4_Metal.pdf") as pdf:
        for page in pdf.pages:
            if "TABLE 4.15" not in (page.extract_text() or ""):
                continue
            for tb in (page.extract_tables() or []):
                for row in tb:
                    cells = [(c or "") for c in row]
                    if cells and cells[0].replace("\n", " ").strip().startswith("CWPB"):
                        techs = cells[0].split("\n")
                        cf4 = [float(x) for x in cells[1].split("\n")]
                        c2f6 = [float(x) for x in cells[3].split("\n")]
                        return {t.strip(): (cf4[i], c2f6[i])
                                for i, t in enumerate(techs)}
    raise RuntimeError("IPCC Table 4.15 PFC defaults not found")


def parse_ar6_gwp() -> dict:
    """AR6 GWP-100.

    CH4/N2O come from Ch.7 Table 7.15 (values written as 'X +/- Y'; GWP-100 is
    the 4th such group).  Halocarbons come from Supplementary Table 7.SM.7,
    where GWP-100 is the 6th numeric token after the species/formula
    (lifetime, RE, AGWP-20, GWP-20, AGWP-100, GWP-100).
    """
    out = {}
    with pdfplumber.open(RAW / "IPCC_AR6_WGI_Chapter07.pdf") as pdf:
        txt = "\n".join((pdf.pages[i].extract_text() or "") for i in range(90, 100))
    for species, key in (("CH4-fossil", "CH4_fossil"),
                         ("CH4-non fossil", "CH4_nonfossil"),
                         ("N2O", "N2O")):
        m = re.search(re.escape(species) + r"\s+(.+)", txt)
        if not m:
            raise RuntimeError(f"AR6 Table 7.15 species '{species}' not found")
        groups = re.findall(r"([\d,]+\.?\d*)\s*±", m.group(1))
        if len(groups) < 4:
            raise RuntimeError(f"AR6 Table 7.15 '{species}': too few +/- groups")
        out[key] = float(groups[3].replace(",", ""))

    with pdfplumber.open(RAW / "IPCC_AR6_WGI_Chapter07_SM.pdf") as pdf:
        smtxt = "\n".join((pdf.pages[i].extract_text() or "") for i in range(14, 24))
    for species in ("HCFC-22", "HFC-32", "HFC-134a", "HFC-125", "HFC-143a",
                    "PFC-14", "PFC-116", "SF6"):
        m = re.search(r"^" + re.escape(species) + r"\s+(.*)$", smtxt, re.M)
        if not m:
            raise RuntimeError(f"AR6 Table 7.SM.7 species '{species}' not found")
        toks = m.group(1).split()
        nums = []
        for t in toks:
            t2 = t.replace(",", "")
            try:
                nums.append(float(t2))
            except ValueError:
                if nums:
                    break  # formula token only ever leads
        if len(nums) < 6:
            raise RuntimeError(f"AR6 Table 7.SM.7 '{species}': too few numbers")
        out[species] = nums[5]
    return out


# EPA SNAP column labels -> AR6 Table 7.SM.7 species names. The EPA table groups
# columns under HCFCs / HFCs / HCs / HFOs / PFCs, so the bare numeric labels are
# disambiguated by that group: 22/124/142b are HCFCs, the rest of the numerics HFCs.
EPA_TO_AR6 = {
    "22": "HCFC-22", "124": "HCFC-124", "142b": "HCFC-142b",
    "23": "HFC-23", "32": "HFC-32", "125": "HFC-125", "134a": "HFC-134a",
    "143a": "HFC-143a", "152a": "HFC-152a", "227ea": "HFC-227ea",
}


def parse_epa_blend_compositions() -> dict:
    """US EPA SNAP 'Compositions of Refrigerant Blends'.

    Two-level header: a group row (HCFCs/HFCs/HCs/HFOs/PFCs) then a substance
    row of 20 labels. Data rows carry 22 cells -- trade name, ASHRAE number,
    then the 20 substance percentages positionally.

    Returns {'R-410A': {'HFC-32': 0.5, 'HFC-125': 0.5}, ...} in mass fractions.
    """
    import html as _html
    raw = (RAW / "EPA_SNAP_refrigerant_blend_compositions.html").read_text(
        encoding="utf-8", errors="replace")
    rows = re.findall(r"<tr[^>]*>(.*?)</tr>", raw, re.S)

    def cells(r):
        return [re.sub(r"\s+", " ", _html.unescape(re.sub(r"<[^>]+>", "", c))).strip()
                for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", r, re.S)]

    sub, out = None, {}
    for r in rows:
        c = cells(r)
        if sub is None:
            if len(c) >= 20 and c[0] == "22" and "32" in c and "125" in c:
                sub = c
            continue
        if len(c) < 3 or not c[1].startswith("R-"):
            continue
        comp = {}
        for i, v in enumerate(c[2:]):
            if i >= len(sub) or not v:
                continue
            try:
                pct = float(v)
            except ValueError:
                continue
            comp[sub[i]] = pct / 100.0
        if comp:
            out[c[1]] = comp
    if not out:
        raise RuntimeError("EPA SNAP blend composition table not parsed")
    return out


def epa_blend_to_ar6(blend: str, comps: dict, ar6: dict):
    """Map one EPA blend to AR6 species. Returns (weights, None) or (None, reason)."""
    if blend not in comps:
        return None, f"blend '{blend}' is not listed in the EPA SNAP table"
    raw = comps[blend]
    total = sum(raw.values())
    if abs(total - 1.0) > 0.02:
        return None, (f"EPA SNAP composition for {blend} sums to {total*100:.1f}%, "
                      "not ~100% -- refusing to derive from an incomplete row")
    weights = {}
    for label, frac in raw.items():
        species = EPA_TO_AR6.get(label)
        if species is None:
            return None, (f"EPA component label '{label}' in {blend} has no AR6 "
                          "species mapping in EPA_TO_AR6")
        if species not in ar6:
            return None, (f"AR6 GWP-100 for component {species} of {blend} was not "
                          "parsed from Table 7.SM.7")
        weights[species] = frac
    return weights, None


def parse_defra_r404a_composition() -> dict | None:
    """DEFRA 2026 methodology para 4.6 states the R-404A blend composition."""
    with pdfplumber.open(RAW / "DEFRA_2026_methodology.pdf") as pdf:
        txt = "\n".join((pdf.pages[i].extract_text() or "") for i in range(40, 48))
    m = re.search(r"R404a that comprises of\s*(\d+)%\s*HFC125\d*,?\s*"
                  r"(\d+)%\s*HFC143a\s*and\s*(\d+)%\s*HFC134a", txt)
    if not m:
        return None
    return {"HFC-125": int(m.group(1)) / 100,
            "HFC-143a": int(m.group(2)) / 100,
            "HFC-134a": int(m.group(3)) / 100}


# --------------------------------------------------------------------------
# row assembly
# --------------------------------------------------------------------------
COLUMNS = ["factor_id", "source_class", "activity", "unit", "gas",
           "value_primary", "source_primary", "vintage_primary",
           "value_crosscheck", "source_crosscheck", "vintage_crosscheck",
           "deviation_pct", "flags", "notes"]

ROWS: list[dict] = []


def _num(x, sig=6):
    if x is None or x == "":
        return ""
    if isinstance(x, str):
        return x
    v = float(x)
    if v == 0:
        return "0"
    s = f"{v:.{sig}g}"
    return s


def add(factor_id, source_class, activity, unit, gas, value_primary,
        source_primary, vintage_primary, value_crosscheck=None,
        source_crosscheck="", vintage_crosscheck="", flags="", notes=""):
    dev = ""
    fl = [f for f in flags.split("|") if f]
    if (isinstance(value_primary, (int, float)) and value_primary not in (0, None)
            and isinstance(value_crosscheck, (int, float))):
        d = (float(value_crosscheck) - float(value_primary)) / float(value_primary) * 100.0
        dev = f"{d:.2f}"
        if abs(d) > 5.0 and "deviation-gt-5pct" not in fl:
            fl.append("deviation-gt-5pct")
    if value_primary in (None, "") and "TODO" not in fl:
        fl.append("TODO")
    ROWS.append({
        "factor_id": factor_id, "source_class": source_class,
        "activity": activity, "unit": unit, "gas": gas,
        "value_primary": _num(value_primary), "source_primary": source_primary,
        "vintage_primary": vintage_primary,
        "value_crosscheck": _num(value_crosscheck),
        "source_crosscheck": source_crosscheck,
        "vintage_crosscheck": vintage_crosscheck,
        "deviation_pct": dev, "flags": "|".join(fl), "notes": notes,
    })


# GHGP fuel key -> (id stem, human activity, units, DEFRA cross-check fuel)
FUEL_MAP = [
    ("gas/diesel oil", "diesel", "Diesel / gas oil (stationary + DG sets + vehicles)",
     ["litre", "tonne"], "diesel (100% mineral diesel)"),
    ("motor gasoline", "petrol", "Petrol / motor gasoline",
     ["litre", "tonne"], "petrol (100% mineral petrol)"),
    ("natural gas", "natural_gas", "Natural gas",
     ["m3", "tonne"], "natural gas (100% mineral blend)"),
    ("liquified petroleum gases", "lpg", "LPG",
     ["litre", "tonne"], "lpg"),
    ("residual fuel oil", "fuel_oil", "Fuel oil (residual / furnace oil)",
     ["litre", "tonne"], "fuel oil"),
    ("other bituminous coal", "coal_bituminous", "Bituminous / steam coal",
     ["tonne"], "coal (industrial)"),
    ("lignite", "lignite", "Lignite", ["tonne"], None),
    ("petroleum coke", "pet_coke", "Petroleum coke", ["tonne"], "petroleum coke"),
    ("coke oven coke", "coke_oven_coke", "Coke (coke-oven coke)", ["tonne"], None),
    ("naphtha", "naphtha", "Naphtha", ["litre", "tonne"], "naphtha"),
    ("other primary solid biomass fuels", "biomass_rice_husk",
     "Biomass / rice husk (solid primary biomass)", ["tonne"], None),
]

UNIT_TEXT = {
    "litre": ("kg {gas} per litre", "per_litre", "litres"),
    "tonne": ("kg {gas} per tonne", "per_tonne", "tonnes"),
    "m3": ("kg {gas} per m3", "per_m3", "cubic metres"),
}

DEFRA_NO_XC = {
    "lignite": "DEFRA 2026 publishes no lignite/brown-coal factor; cross-check left empty.",
    "coke_oven_coke": ("DEFRA 2026 lists 'Coking coal' (a kiln INPUT coal), not coke-oven "
                       "coke (the OUTPUT). Not the same commodity, so no cross-check is "
                       "forced. DEFRA coking coal = 3144.16 kgCO2/t for reference only."),
    "biomass_rice_husk": ("DEFRA 'Bioenergy' publishes only the non-CO2 portion as an "
                          "AR5-weighted kgCO2e/tonne aggregate (e.g. Grass/straw 46.89 "
                          "kgCO2e/t), with biogenic CO2 excluded. Cannot be reconciled to "
                          "a per-gas mass factor, so cross-check left empty."),
}


def build_fuels(ghgp, defra, defra_gwp):
    ch4_gwp = defra_gwp["methane"]        # AR5, 28
    n2o_gwp = defra_gwp["nitrous oxide"]  # AR5, 265
    for gkey, stem, activity, units, dfuel in FUEL_MAP:
        f = ghgp.get(gkey)
        if f is None:
            # A miss here is a mapping bug in this script, not a gap in the source,
            # so fail loudly rather than silently degrading a core fuel to TODO.
            raise RuntimeError(
                f"FUEL_MAP key '{gkey}' does not match any GHGP fuel. "
                f"Available: {sorted(ghgp)}")
        biogenic = stem == "biomass_rice_husk"
        for u in units:
            unit_tmpl, col, defra_unit = UNIT_TEXT[u]
            for gas in ("CO2", "CH4", "N2O"):
                val = f[gas][col]
                if not isinstance(val, (int, float)):
                    add(f"fuel_{stem}_{gas.lower()}_per_{u}", "fuel", activity,
                        unit_tmpl.format(gas=gas), gas, None, S_GHGP, V_GHGP,
                        flags="TODO",
                        notes=f"GHGP publishes no {u}-basis value for "
                              f"'{f['label']}' (density not given in the workbook).")
                    continue

                # --- cross-check ---
                xc = xcs = xcv = None
                note_bits = [
                    f"GHGP Stationary Combustion Table "
                    f"{'1' if gas == 'CO2' else '2' if gas == 'CH4' else '3'}, "
                    f"fuel '{f['label']}'; NCV {f['ncv']} TJ/Gg (net/lower heating value basis)."
                ]
                if u == "litre" and f["dens_l"]:
                    note_bits.append(f"GHGP liquid density {f['dens_l']:.5f} kg/litre.")
                if u == "m3" and f["dens_m3"]:
                    note_bits.append(f"GHGP gas density {f['dens_m3']} kg/m3.")

                if dfuel:
                    d = defra.get((dfuel, defra_unit))
                    if d:
                        if gas == "CO2":
                            xc = d["co2"]
                            note_bits.append(
                                f"Cross-check: DEFRA 2026 Fuels, '{dfuel}', {defra_unit}, "
                                "column 'kg CO2e of CO2 per unit' (= kg CO2, GWP-invariant).")
                        else:
                            raw = d["ch4_co2e"] if gas == "CH4" else d["n2o_co2e"]
                            gwp = ch4_gwp if gas == "CH4" else n2o_gwp
                            if isinstance(raw, (int, float)) and raw:
                                xc = raw / gwp
                                note_bits.append(
                                    f"Cross-check unit reconciliation: DEFRA publishes "
                                    f"{gas} pre-weighted as CO2e; divided by DEFRA's own "
                                    f"AR5 GWP ({gwp}) to recover kg {gas} per {defra_unit[:-1]}. "
                                    "GHGP uses IPCC 2006 generic Tier 1 defaults while DEFRA "
                                    "uses UK GHGI measured factors, so a large divergence "
                                    "here is methodological, not an error.")
                        if xc is not None:
                            xcs, xcv = S_DEFRA, V_DEFRA
                    else:
                        note_bits.append(
                            f"No DEFRA 2026 row for ('{dfuel}', '{defra_unit}'); "
                            "cross-check left empty.")
                elif stem in DEFRA_NO_XC:
                    note_bits.append(DEFRA_NO_XC[stem])

                flags = ["tier1-default"]
                if biogenic and gas == "CO2":
                    flags.append("biogenic-memo")
                    note_bits.append(
                        "BIOGENIC: report as a memo item, NOT in the Scope 1 total "
                        "(SPEC 5, textiles husk boilers). CH4 and N2O from the same fuel "
                        "ARE normal Scope 1 lines.")

                add(f"fuel_{stem}_{gas.lower()}_per_{u}", "fuel", activity,
                    unit_tmpl.format(gas=gas), gas, val, S_GHGP, V_GHGP,
                    xc, xcs or "", xcv or "", flags="|".join(flags),
                    notes=" ".join(note_bits))


def build_grid(cea):
    add("grid_india_weighted_avg", "grid",
        "India national grid electricity (location-based Scope 2)",
        "tCO2 per MWh", "CO2", cea["avg"], S_CEA, V_CEA,
        cea["prev_fy_avg"], S_CEA + " Annexure-I (prior vintage)",
        "FY 2023-24 (as published in v21.0)",
        flags="",
        notes="CEA v21.0 User Guide Table S, page 1: weighted average emission factor of "
              "the unified Indian Grid for FY 2024-25, adjusted for cross-border transfers, "
              "including RES and grid-connected captive injection. Cross-check is the "
              "FY 2023-24 value from Annexure-I of the SAME document (0.727), i.e. the "
              "prior-vintage reference the SPEC calls for; the deviation is the real "
              "year-on-year decarbonisation of the grid, not a source disagreement. "
              "Unadjusted (excluding cross-border transfers) value is 0.712.")
    for key, name, desc in (
        ("om", "grid_india_operating_margin", "simple operating margin (OM)"),
        ("bm", "grid_india_build_margin", "build margin (BM)"),
        ("cm", "grid_india_combined_margin", "combined margin (CM, 50:50 OM/BM)"),
    ):
        add(name, "grid", f"India national grid electricity - {desc}",
            "tCO2 per MWh", "CO2", cea[key], S_CEA, V_CEA,
            flags="",
            notes=f"CEA v21.0 User Guide Table S, page 1, {desc}, FY 2024-25, adjusted for "
                  "cross-border transfers. Not used for location-based Scope 2; carried for "
                  "CDM/Article-6 style marginal calculations and for the market-based line.")


GWP_SPEC = [
    ("gwp_ch4_fossil", "Methane (fossil origin)", "CH4", "CH4_fossil", "methane",
     "AR6 separates fossil from non-fossil CH4; fossil is the correct choice for "
     "natural gas, LPG, coal and oil combustion."),
    ("gwp_ch4_nonfossil", "Methane (non-fossil / biogenic origin)", "CH4",
     "CH4_nonfossil", "methane",
     "Use for CH4 from biomass/rice-husk combustion and effluent treatment."),
    ("gwp_n2o", "Nitrous oxide", "N2O", "N2O", "nitrous oxide", ""),
    ("gwp_cf4", "Tetrafluoromethane CF4 (PFC-14)", "CF4", "PFC-14",
     "perfluoromethane (pfc-14)", "Aluminium smelting PFC."),
    ("gwp_c2f6", "Hexafluoroethane C2F6 (PFC-116)", "C2F6", "PFC-116",
     "perfluoroethane (pfc-116)", "Aluminium smelting PFC."),
    ("gwp_sf6", "Sulphur hexafluoride", "SF6", "SF6",
     "sulphur hexafluoride (sf6)", "Switchgear; carried for completeness."),
    ("gwp_r22", "Refrigerant R-22 (HCFC-22)", "R-22", "HCFC-22",
     "hcfc-22/r22 = chlorodifluoromethane",
     "Montreal-Protocol gas: still widely charged in Indian industrial chillers. "
     "Not a Kyoto basket gas, so DEFRA reports it under 'non-Kyoto'."),
    ("gwp_r32", "Refrigerant R-32 (HFC-32)", "R-32", "HFC-32", "hfc-32", ""),
    ("gwp_r134a", "Refrigerant R-134a (HFC-134a)", "R-134a", "HFC-134a",
     "hfc-134a", ""),
]


def build_gwp(ar6, defra_gwp, epa_comps):
    for fid, activity, gas, ar6key, defra_key, extra in GWP_SPEC:
        v = ar6.get(ar6key)
        xc = defra_gwp.get(defra_key)
        src = S_AR6 if ar6key in ("CH4_fossil", "CH4_nonfossil", "N2O") else S_AR6SM
        note = (f"AR6 GWP-100, species '{ar6key}'. DECLARED BASIS FOR THIS PROJECT = AR6. "
                "Cross-check is DEFRA 2026, which is an AR5-basis set (its methodology "
                "paper para 1 and 4.5 state GWP CH4=28, N2O=265 from IPCC AR5); the "
                "deviation column here therefore measures the AR5->AR6 revision, NOT a "
                "source error. Do not mix the two sets. " + extra).strip()
        add(fid, "gwp", activity, "kg CO2e per kg gas", gas, v, src, V_AR6,
            xc, S_DEFRA + " (AR5 basis)", V_DEFRA,
            flags="ar5-vs-ar6", notes=note)

    # --- Refrigerant blends. AR6 assesses pure species only and publishes no blend
    # GWPs, so these are DERIVED as sum(w_i x GWP_i) from AR6 component values and
    # the EPA SNAP published mass fractions. Flagged derived-blend, never guessed.
    for fid, blend, extra in (
        ("gwp_r404a", "R-404A",
         "Cross-confirmation: the DEFRA 2026 methodology paper para 4.6 independently "
         "states the same split (44% HFC-125, 52% HFC-143a, 4% HFC-134a) and shows its "
         "own AR5 arithmetic [3170 x 0.44] + [4800 x 0.52] + [1300 x 0.04] = 3943, which "
         "is the cross-check value here. Two independent sources agree on the "
         "composition."),
        ("gwp_r410a", "R-410A",
         "Cross-confirmation: DEFRA's AR5-basis published value for R-410A is 1924 "
         "kgCO2e/kg (the cross-check here); applying DEFRA's own AR5 component GWPs "
         "(HFC-32 = 677, HFC-125 = 3170) to this same 50/50 split reproduces "
         "0.5 x 677 + 0.5 x 3170 = 1923.5, i.e. DEFRA's published 1924 to within "
         "rounding. That reproduction independently validates the 50/50 composition."),
    ):
        xc = defra_gwp.get(blend.lower().replace("-", ""))
        weights, reason = epa_blend_to_ar6(blend, epa_comps, ar6)
        if weights is None:
            add(fid, "gwp", f"Refrigerant {blend} (blend)", "kg CO2e per kg gas",
                blend, None, "", V_AR6, xc, S_DEFRA + " (AR5 basis)", V_DEFRA,
                flags="TODO",
                notes=f"MISSING: AR6-basis GWP for {blend}. AR6 assesses pure species only, "
                      f"so a blend value must be derived from composition. Blocked because: "
                      f"{reason}. The DEFRA AR5 value in value_crosscheck is context ONLY "
                      "and must not be promoted without changing the declared GWP basis.")
            continue
        val = sum(ar6[sp] * w for sp, w in weights.items())
        comp_txt = ", ".join(
            f"{w*100:g}% {sp} (AR6 GWP-100 {ar6[sp]:g})"
            for sp, w in sorted(weights.items(), key=lambda kv: -kv[1]))
        calc_txt = " + ".join(
            f"{w:g} x {ar6[sp]:g}"
            for sp, w in sorted(weights.items(), key=lambda kv: -kv[1]))
        add(fid, "gwp", f"Refrigerant {blend} (blend)", "kg CO2e per kg gas",
            blend, val,
            S_AR6SM + " components x " + S_EPA_SNAP + " mass fractions", V_AR6,
            xc, S_DEFRA + " (AR5 basis)", V_DEFRA,
            flags="ar5-vs-ar6|derived-blend",
            notes=f"DERIVED, not published: AR6 assesses pure species only and publishes no "
                  f"blend GWPs. METHOD: mass-weighted sum of constituent GWP-100 values, the "
                  f"standard convention (EU F-gas Regulation 517/2014 Annex IV and (EU) "
                  f"2024/573 Annex VI both define blend GWP exactly this way; both were "
                  f"fetched and confirmed to give the method but no composition table). "
                  f"COMPOSITION by weight: {comp_txt}, parsed from the US EPA SNAP "
                  f"'Compositions of Refrigerant Blends' table. CALCULATION: {calc_txt} = "
                  f"{val:g} kgCO2e/kg. {extra} The deviation against the cross-check is "
                  "therefore purely the AR5->AR6 component revision, not a disagreement "
                  "about composition.")


def build_process(cem, chem, met):
    add("process_cement_clinker_co2", "process",
        "Cement clinker calcination (CO2 from carbonate decomposition)",
        "tCO2 per tonne clinker", "CO2", cem["clinker_tier1"], S_IPCC2, V_IPCC,
        cem["clinker_base"], S_IPCC2 + " (uncorrected base EF)", V_IPCC,
        flags="tier1-default",
        notes="IPCC 2006 Vol.3 Ch.2 Equation 2.4 (p.2.12): EF_clc = 0.51 x 1.02 (CKD "
              "correction) = 0.52 tCO2/t clinker. Assumes clinker is 65% CaO, 100% derived "
              "from CaCO3, 100% calcination. Cross-check is the same equation's uncorrected "
              "base factor 0.51 (the Tier 2 starting point, which excludes cement kiln dust). "
              "SPEC 5 prefers a plant-specific CaO-based Tier 2 factor where available: "
              "Ch.2 gives 0.47 for 60% CaO and 0.53 for 67% CaO.")

    for fid, key, unit, activity, note in (
        ("process_steel_eaf_tier1", "steel_eaf", "tCO2 per tonne crude steel",
         "Steel - Electric Arc Furnace (EAF) process CO2",
         "Process CO2 only (electrode + charge carbon + flux). Excludes the grid or captive "
         "electricity that dominates an EAF footprint - that is the Scope 2 / captive Scope 1 "
         "line, not this factor."),
        ("process_steel_bof_tier1", "steel_bof", "tCO2 per tonne crude steel",
         "Steel - Basic Oxygen Furnace (BOF) process CO2", "For the BF-BOF route gate."),
        ("process_steel_bf_iron_tier1", "steel_bf_iron", "tCO2 per tonne pig iron",
         "Steel - blast furnace pig iron production", "For the BF-BOF route gate."),
        ("process_steel_dri_tier1", "steel_dri", "tCO2 per tonne DRI",
         "Steel - direct reduced iron (DRI) production",
         "IPCC Tier 1 assumes NATURAL-GAS-based DRI (12.5 GJ/t, 15.3 kgC/GJ). India's "
         "coal-DRI route is materially more carbon-intensive, so SPEC 5's coal-DRI gate must "
         "NOT use this default - compute from actual reductant instead."),
        ("process_steel_coke_oven_tier1", "steel_coke_oven", "tCO2 per tonne coke",
         "Steel - coke oven (integrated coke production)", ""),
        ("process_steel_sinter_tier1", "steel_sinter", "tCO2 per tonne sinter",
         "Steel - sinter production", ""),
    ):
        add(fid, "process", activity, unit, "CO2", met[key], S_IPCC4, V_IPCC,
            flags="tier1-default",
            notes=("IPCC 2006 Vol.3 Ch.4 Table 4.1 (Tier 1 default CO2 emission factors for "
                   "coke production and iron & steel production). " + note).strip())

    add("process_eaf_electrode_carbon", "process",
        "Steel - EAF carbon electrode consumption (carbon content)",
        "kg C per kg electrode", "C", met["eaf_electrode_c"], S_IPCC4, V_IPCC,
        flags="tier1-default",
        notes="IPCC 2006 Vol.3 Ch.4 Table 4.3 (Tier 2 material-specific carbon contents), row "
              "'EAF Carbon Electrodes' = 0.82 kg C/kg, assumed 80% petroleum coke / 20% coal "
              "tar. This is a CARBON CONTENT, not a CO2 factor - see the derived row "
              "process_eaf_electrode_co2 for the CO2 form.")
    co2_e = met["eaf_electrode_c"] * 44.0 / 12.0
    add("process_eaf_electrode_co2", "process",
        "Steel - EAF carbon electrode consumption (CO2)",
        "tCO2 per tonne electrode", "CO2", co2_e, S_IPCC4 + " x 44/12", V_IPCC,
        flags="tier1-default|stoichiometric|derived",
        notes=f"DERIVED: {met['eaf_electrode_c']} kg C/kg electrode (IPCC Vol.3 Ch.4 Table 4.3) "
              f"x 44/12 = {co2_e:.5f} tCO2 per tonne electrode. The 44/12 ratio is stated "
              "verbatim in the same chapter (Ch.4, Equation 4.24 variable list: "
              "'44/12 = CO2 molecular mass : carbon atomic mass ratio'). Assumes complete "
              "oxidation of consumed electrode carbon.")

    for fid, key, tech in (("process_alu_anode_prebake_co2", "al_anode_prebake", "Prebake"),
                           ("process_alu_paste_soderberg_co2", "al_anode_soderberg",
                            "Soderberg")):
        add(fid, "process",
            f"Aluminium - {tech} anode/paste consumption CO2",
            "tCO2 per tonne Al", "CO2", met[key], S_IPCC4, V_IPCC,
            flags="tier1-default",
            notes=f"IPCC 2006 Vol.3 Ch.4 Table 4.10 (Tier 1 technology-specific EFs for "
                  f"calculating CO2 from anode or paste consumption), {tech} = {met[key]} "
                  "tCO2/t Al, uncertainty +/-10%. Derived from International Aluminium "
                  "Institute Life Cycle Assessment of Aluminium (IAI, 2000). The Prebake "
                  "factor already includes CO2 from combustion of pitch volatiles and packing "
                  "coke during anode baking. SPEC 5 prefers the Tier 2/3 net-anode-consumption "
                  "route (anode C x 44/12) where plant data exist.")

    for tech, (cf4, c2f6) in met["pfc"].items():
        for gas, val in (("CF4", cf4), ("C2F6", c2f6)):
            add(f"process_alu_{gas.lower()}_{tech.lower()}", "process",
                f"Aluminium - PFC from anode effects, {tech} cell technology",
                "kg gas per tonne Al", gas, val, S_IPCC4, V_IPCC,
                flags="tier1-default",
                notes=f"IPCC 2006 Vol.3 Ch.4 Table 4.15 (Tier 1 default PFC EFs by cell "
                      f"technology), {tech} {gas} = {val} kg/t Al. Technology codes: CWPB = "
                      "Centre Worked Prebake, SWPB = Side Worked Prebake, VSS = Vertical Stud "
                      "Soderberg, HSS = Horizontal Stud Soderberg. Uncertainty is very wide "
                      "(CWPB -99%/+380%). IPCC states these defaults 'should only be used in "
                      "the absence of Tier 2 or Tier 3 data' - SPEC 5 therefore requires the "
                      "anode-effect slope method where AE data exist, and this row must carry "
                      "the 'default' flag into the calculation trail when used.")

    for fid, key, activity, abated in (
        ("process_hno3_n2o_nscr", "hno3_nscr",
         "Nitric acid - plants with NSCR abatement (all processes)", True),
        ("process_hno3_n2o_destruction", "hno3_destruction",
         "Nitric acid - plants with process-integrated or tailgas N2O destruction", True),
        ("process_hno3_n2o_atmospheric", "hno3_atmospheric",
         "Nitric acid - atmospheric (low) pressure plants, unabated", False),
        ("process_hno3_n2o_medium", "hno3_medium",
         "Nitric acid - medium pressure combustion plants, unabated", False),
        ("process_hno3_n2o_high", "hno3_high",
         "Nitric acid - high pressure plants, unabated", False),
    ):
        unc = {"hno3_nscr": "+/-10%", "hno3_destruction": "+/-10%",
               "hno3_atmospheric": "+/-10%", "hno3_medium": "+/-20%",
               "hno3_high": "+/-40%"}[key]
        add(fid, "process", activity, "kg N2O per tonne nitric acid (100% HNO3)",
            "N2O", chem[key], S_IPCC3, V_IPCC, flags="tier1-default",
            notes=f"IPCC 2006 Vol.3 Ch.3 Table 3.3 (default factors for nitric acid "
                  f"production), {unc}. Source: van Balken (2005). Factors relate to 100 "
                  "percent pure acid - normalise plant tonnage to 100% HNO3 first. "
                  f"{'ABATED' if abated else 'UNABATED'} case. The NSCR and destruction "
                  "factors already incorporate the abatement effect, so do NOT apply a "
                  "separate destruction efficiency on top; IPCC requires verifying the "
                  "abatement kit is installed AND operated year-round before using them. "
                  "Convert to CO2e with the AR6 N2O GWP row gwp_n2o (SPEC 5 fertilisers line).")


def build_stoichiometric(cem, chem):
    add("stoich_limestone_caco3", "stoichiometric",
        "Limestone flux calcination (CaCO3 -> CaO + CO2)",
        "tCO2 per tonne CaCO3", "CO2", cem["caco3"],
        S_IPCC2 + " Table 2.1 (= molar mass ratio)", V_IPCC,
        0.440, "SPEC 4 pinned constant / molar mass", "exact chemistry",
        flags="stoichiometric",
        notes="IPCC 2006 Vol.3 Ch.2 Table 2.1, Calcite/aragonite CaCO3, formula weight "
              "100.0869, EF 0.43971 tCO2 per tonne carbonate at 100% calcination. This is a "
              "molar-mass identity (44.0095/100.0869), not an empirical estimate. SPEC 4 pins "
              "the rounded 0.440 - recorded as cross-check; the 0.07% difference is rounding "
              "only. Used for steel/cement flux (SPEC 5 steel route gate).")
    add("stoich_dolomite", "stoichiometric",
        "Dolomite flux calcination (CaMg(CO3)2 -> CaO + MgO + 2 CO2)",
        "tCO2 per tonne dolomite", "CO2", cem["dolomite"],
        S_IPCC2 + " Table 2.1 (= molar mass ratio)", V_IPCC,
        0.477, "SPEC 4 pinned constant / molar mass", "exact chemistry",
        flags="stoichiometric",
        notes="IPCC 2006 Vol.3 Ch.2 Table 2.1, Dolomite CaMg(CO3)2, formula weight 184.4008, "
              "EF 0.47732 tCO2 per tonne carbonate at 100% calcination (2 mol CO2 per mol "
              "dolomite). Molar-mass identity. SPEC 4 pins the rounded 0.477 - recorded as "
              "cross-check.")
    add("stoich_urea_co2_uptake", "stoichiometric",
        "Urea production - CO2 consumed (negative / uptake line)",
        "tCO2 per tonne urea", "CO2", chem["urea_uptake"],
        S_IPCC3 + " (urea production, p.3.16)", V_IPCC,
        0.733, "SPEC 4 pinned constant / molar mass", "exact chemistry",
        flags="stoichiometric",
        notes="IPCC 2006 Vol.3 Ch.3, Urea Production section, stated verbatim: 'Assuming "
              "complete conversion of NH3 and CO2 to urea, 0.733 tonnes of CO2 are required "
              "per tonne of urea produced.' SIGN CONVENTION: this is CO2 CONSUMED, so SPEC 5 "
              "requires it as an explicit NEGATIVE line (-0.733) in the fertiliser module, "
              "netted against the ammonia-plant feedstock CO2 - never silently absorbed. The "
              "CO2 is re-released downstream when the urea is applied (Vol.4 agriculture), "
              "which is outside this project's Scope 1+2 boundary.")
    c_ratio = 44.0 / 12.0
    add("stoich_carbon_to_co2", "stoichiometric",
        "Carbon to CO2 conversion ratio", "tCO2 per tonne C", "CO2", c_ratio,
        S_IPCC4 + " (Eq. 4.24 variable list) / molar mass", V_IPCC,
        3.667, "SPEC 4 pinned constant", "exact chemistry",
        flags="stoichiometric",
        notes="44/12 = 3.66667. Stated verbatim in IPCC 2006 Vol.3 Ch.4 (Equation 4.24 "
              "variable list): '44/12 = CO2 molecular mass : carbon atomic mass ratio'. Used "
              "for aluminium anode carbon, EAF electrode carbon, and fertiliser feedstock "
              "carbon (SPEC 5).")


def build_steam():
    add("steam_purchased", "steam", "Purchased steam / heat (Scope 2)",
        "tCO2e per tonne steam", "CO2e", None, "", "",
        None, S_DEFRA + " 'Heat and steam' (UK grid context only)", V_DEFRA,
        flags="TODO",
        notes="NO FACTOR ROW BY DESIGN (SPEC 5, line C6). Purchased steam is supplier-specific: "
              "the emission factor depends on the supplier's boiler fuel, boiler efficiency and "
              "any cogeneration allocation, none of which are properties of a factor library. "
              "REQUIRED ENGINE BEHAVIOUR: (1) prefer a supplier-documented tCO2e per tonne (or "
              "per GJ) of steam, recorded with its own source+vintage in the calculation trail; "
              "(2) fall back to a documented boiler-efficiency reconstruction - steam enthalpy "
              "demand / boiler efficiency = fuel energy, then apply the matching fuel_* rows "
              "from this file - with the assumed efficiency surfaced as an input and the line "
              "flagged 'estimate'; (3) never substitute a generic factor silently. DEFRA 2026 "
              "'Heat and steam' publishes 0.17529 kgCO2e/kWh (onsite and district, 2026) but "
              "that is a UK-supply-mix number and is NOT applicable to Indian suppliers - it is "
              "recorded here as context only, not as a usable default. The boiler-efficiency "
              "fallback must be defined in engine/common.py before this line can run.")


# --------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--refresh", action="store_true",
                    help="re-download all sources even if cached")
    args = ap.parse_args()

    print(f"build_factors.py  (sources retrieved {RETRIEVED})")
    print("1. fetching sources into data/raw/ ...")
    fetch_all(refresh=args.refresh)

    print("2. parsing ...")
    ghgp = parse_ghgp_fuels();            print(f"   GHGP fuels          : {len(ghgp)}")
    defra = parse_defra_fuels();          print(f"   DEFRA fuel rows     : {len(defra)}")
    defra_gwp = parse_defra_gwp();        print(f"   DEFRA GWP rows      : {len(defra_gwp)}")
    cea = parse_cea();                    print(f"   CEA FY2024-25       : {cea}")
    cem = parse_ipcc_clinker();           print(f"   IPCC Ch2            : {cem}")
    chem = parse_ipcc_chemical();         print(f"   IPCC Ch3            : {chem}")
    met = parse_ipcc_metal();             print(f"   IPCC Ch4 PFC        : {met['pfc']}")
    ar6 = parse_ar6_gwp();                print(f"   AR6 GWP-100         : {ar6}")
    epa_comps = parse_epa_blend_compositions()
    print(f"   EPA blends parsed   : {len(epa_comps)}")
    for b in ("R-404A", "R-410A"):
        print(f"     {b}: {epa_comps.get(b)}")
    # DEFRA's own R-404A composition, kept as an independent cross-confirmation
    print(f"   DEFRA R-404A comp   : {parse_defra_r404a_composition()}")

    print("3. building rows ...")
    build_fuels(ghgp, defra, defra_gwp)
    build_grid(cea)
    build_gwp(ar6, defra_gwp, epa_comps)
    build_process(cem, chem, met)
    build_stoichiometric(cem, chem)
    build_steam()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS, quoting=csv.QUOTE_MINIMAL)
        w.writeheader()
        w.writerows(ROWS)

    ids = [r["factor_id"] for r in ROWS]
    assert len(ids) == len(set(ids)), "duplicate factor_id"

    by_class: dict[str, int] = {}
    for r in ROWS:
        by_class[r["source_class"]] = by_class.get(r["source_class"], 0) + 1
    todo = [r for r in ROWS if "TODO" in r["flags"]]
    devs = [r for r in ROWS if "deviation-gt-5pct" in r["flags"]]

    print(f"\nwrote {OUT.relative_to(ROOT)}  ({len(ROWS)} rows)")
    for k in sorted(by_class):
        print(f"   {k:<16} {by_class[k]}")
    print(f"   TODO rows        {len(todo)}")
    print(f"   >5% deviations   {len(devs)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
